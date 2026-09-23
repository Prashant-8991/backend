"""
Import observation.xlsx (tag_number, suggestion, final_observation) into cattle_data.
Matches on tag_number and updates the suggestion + final_observation columns.
Source: C:\\Users\\CMO\\Desktop\\data\\Cow Data (07-06-2026)\\COW all data Single sheet\\observation.xlsx
Target: gaushala DB (postgres/postgres@localhost:5433/gaushala)
"""
import openpyxl
import psycopg2
from pathlib import Path

EXCEL_PATH = Path(r"C:\Users\CMO\Desktop\data\Cow Data (07-06-2026)\COW all data Single sheet\observation.xlsx")
DB_CONFIG = {
    "host": "localhost",
    "port": 5433,
    "dbname": "gaushala",
    "user": "postgres",
    "password": "postgres",
}


def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def load_rows():
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb["data"] if "data" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if not r or all(c is None for c in r):
            continue
        tag = clean(r[0]) if len(r) > 0 else None
        suggestion = clean(r[1]) if len(r) > 1 else None
        final_obs = clean(r[2]) if len(r) > 2 else None
        if not tag:
            continue
        rows.append({"tag_number": tag, "suggestion": suggestion, "final_observation": final_obs})
    wb.close()
    return rows


def main():
    print(f"Excel exists: {EXCEL_PATH.exists()} -> {EXCEL_PATH}")
    if not EXCEL_PATH.exists():
        raise SystemExit(f"File not found: {EXCEL_PATH}")
    rows = load_rows()
    print(f"Parsed {len(rows)} rows from Excel")

    conn = psycopg2.connect(**DB_CONFIG)
    cur = conn.cursor()

    updated = 0
    not_found = []
    for rec in rows:
        cur.execute(
            "UPDATE cattle_data SET suggestion = %s, final_observation = %s WHERE tag_number = %s",
            (rec["suggestion"], rec["final_observation"], rec["tag_number"]),
        )
        if cur.rowcount > 0:
            updated += cur.rowcount
        else:
            not_found.append(rec["tag_number"])
    conn.commit()

    print(f"Updated rows: {updated}")
    print(f"Tags not found in cattle_data: {len(not_found)}")
    if not_found:
        print("  ", ", ".join(not_found[:40]))

    # Verification
    cur.execute("SELECT COUNT(*) FROM cattle_data WHERE suggestion IS NOT NULL OR final_observation IS NOT NULL")
    print("Rows with suggestion/final_observation in DB:", cur.fetchone()[0])
    cur.execute("SELECT tag_number, suggestion, final_observation FROM cattle_data WHERE tag_number IN (SELECT tag_number FROM cattle_data WHERE suggestion IS NOT NULL OR final_observation IS NOT NULL LIMIT 5)")
    for row in cur.fetchall():
        print("  ", row)
    cur.close()
    conn.close()


if __name__ == "__main__":
    main()
